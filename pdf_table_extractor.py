#!/usr/bin/env python3
"""
PDF Table Extractor Tool - Enhanced Version
Supports: Tabula, Camelot, PDFPlumber
No virtual environment required
"""

import os
import sys
import argparse
import pandas as pd
from openpyxl import Workbook

# ---------------- Dependency Check ----------------
def check_and_install_dependencies():
    """Check for dependencies and guide installation"""
    missing_packages = []
    modules = {
        "pandas": "pandas",
        "openpyxl": "openpyxl",
        "tabula": "tabula-py",
        "camelot": "camelot-py",
        "pdfplumber": "pdfplumber"
    }

    for mod, pkg in modules.items():
        try:
            __import__(mod)
            print(f"✓ {pkg} available")
        except ImportError:
            missing_packages.append(pkg)

    if missing_packages:
        print(f"\n❌ Missing packages: {', '.join(missing_packages)}")
        print("\nTo install, run one of these commands:")
        print(f"1. pip3 install --user {' '.join(missing_packages)}")
        print(f"2. Or install individually: pip3 install --user {missing_packages[0]}")
        sys.exit(1)
    
    return True


# ---------------- Table Extractor Class ----------------
class SimplePDFTableExtractor:
    def __init__(self, method="auto"):
        self.tables = []
        self.method = method.lower()

    def extract_tables(self, pdf_path):
        """Extract tables using chosen or fallback methods"""
        print(f"\n📄 Processing PDF: {pdf_path}")
        all_tables = []

        # Method 1: Tabula
        if self.method in ("auto", "tabula"):
            try:
                import tabula
                print("🔹 Trying extraction with Tabula...")
                tabula_tables = tabula.read_pdf(
                    pdf_path, pages='all',
                    multiple_tables=True,
                    lattice=True, stream=True
                )
                for i, table in enumerate(tabula_tables):
                    if not table.empty:
                        table = table.dropna(how='all').dropna(axis=1, how='all')
                        if len(table) > 0:
                            all_tables.append(table)
                            print(f"  - Tabula found table {i+1}: {table.shape[1]}x{table.shape[0]}")
            except Exception as e:
                print(f"  ⚠️ Tabula error: {e}")
                if self.method == "tabula":
                    return []

        # Method 2: Camelot
        if (self.method in ("auto", "camelot")) and not all_tables:
            try:
                import camelot
                print("🔹 Trying extraction with Camelot...")
                camelot_tables = camelot.read_pdf(pdf_path, pages='all')
                for i, table in enumerate(camelot_tables):
                    df = table.df.dropna(how='all').dropna(axis=1, how='all')
                    if len(df) > 1:
                        all_tables.append(df)
                        print(f"  - Camelot found table {i+1}: {df.shape[1]}x{df.shape[0]}")
            except Exception as e:
                print(f"  ⚠️ Camelot error: {e}")
                if self.method == "camelot":
                    return []

        # Method 3: PDFPlumber
        if (self.method in ("auto", "pdfplumber")) and not all_tables:
            try:
                import pdfplumber
                print("🔹 Trying extraction with PDFPlumber...")
                with pdfplumber.open(pdf_path) as pdf:
                    for page_num, page in enumerate(pdf.pages):
                        tables = page.extract_tables()
                        for table_num, table_data in enumerate(tables):
                            if table_data and len(table_data) > 1:
                                df = pd.DataFrame(table_data)
                                df = df.dropna(how='all').dropna(axis=1, how='all')
                                if len(df) > 1:
                                    all_tables.append(df)
                                    print(f"  - PDFPlumber found table on page {page_num+1}: {df.shape[1]}x{df.shape[0]}")
            except Exception as e:
                print(f"  ⚠️ PDFPlumber error: {e}")

        self.tables = all_tables
        return all_tables

    def save_to_excel(self, output_path):
        """Save all tables to Excel with separate sheets"""
        if not self.tables:
            print("❌ No tables to save!")
            return False
        
        try:
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                for i, table in enumerate(self.tables):
                    sheet_name = f"Table_{i+1}"[:31]
                    table = table.dropna(how='all').dropna(axis=1, how='all').reset_index(drop=True)
                    table.to_excel(writer, sheet_name=sheet_name, index=False)

                    # Auto-fit column width
                    worksheet = writer.sheets[sheet_name]
                    for column in worksheet.columns:
                        max_length = 0
                        column_letter = column[0].column_letter
                        for cell in column:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
            
            print(f"✅ Successfully saved {len(self.tables)} tables to: {output_path}")
            return True
            
        except Exception as e:
            print(f"❌ Error saving Excel file: {e}")
            return False


# ---------------- Main Logic ----------------
def main():
    print("\nPDF Table Extractor")
    print("=" * 60)

    check_and_install_dependencies()

    # Command-line arguments
    parser = argparse.ArgumentParser(description="Extract tables from PDF into Excel.")
    parser.add_argument("pdf", help="Path to the PDF file")
    parser.add_argument("-o", "--output", help="Output Excel file name (default: <pdf>_tables.xlsx)")
    parser.add_argument("-m", "--method", choices=["auto", "tabula", "camelot", "pdfplumber"],
                        default="auto", help="Extraction method (default: auto)")
    args = parser.parse_args()

    pdf_path = os.path.expanduser(args.pdf)
    if not os.path.exists(pdf_path):
        print(f"❌ File not found: {pdf_path}")
        sys.exit(1)

    output_path = args.output or f"{os.path.splitext(pdf_path)[0]}_tables.xlsx"

    # Extraction
    extractor = SimplePDFTableExtractor(method=args.method)
    tables = extractor.extract_tables(pdf_path)

    if not tables:
        print("❌ No tables found in the PDF!")
        print("Try another method: --method=camelot or --method=pdfplumber")
        sys.exit(1)

    print(f"\n✅ Found {len(tables)} tables total")

    # Save results
    extractor.save_to_excel(output_path)

    # Summary
    print("\n📊 Extraction Summary:")
    for i, table in enumerate(tables):
        print(f"  Table {i+1}: {table.shape[1]} columns × {table.shape[0]} rows")
#asasdas

if __name__ == "__main__":
    main()

